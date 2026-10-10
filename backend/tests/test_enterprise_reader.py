import hashlib
import json
import os
from dataclasses import replace
from pathlib import Path
from docx import Document
from app.enterprise_demo import add_enterprise
from app.main import create_app
from app.parser_worker import extract
from app.security import execute
from conftest import ask, client_for


def test_docx_tables_remain_between_paragraphs(tmp_path):
    doc = Document()
    doc.add_paragraph("Before 🚀")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Fictional council"
    doc.add_paragraph("After: exact entitlement")
    path = tmp_path / "ordered.docx"
    doc.save(path)
    result = extract(path)
    assert result["text"] == "Before 🚀\nOwner | Fictional council\nAfter: exact entitlement\n"
    assert [s["locator"] for s in result["segments"]] == [
        "DOCX paragraph 1",
        "DOCX table 1 row 1",
        "DOCX paragraph 2",
    ]


def test_rich_corpus_versions_exact_spans_originals_conflicts_and_revocation(system):
    app, _ = system
    add_enterprise(app.state.engine, app.state.settings)
    reviewer = client_for(system, "ravi")
    employee = client_for(system)
    root = Path(__file__).resolve().parents[2] / "docs/demo-enterprise"
    manifest = json.loads((root / "manifest.json").read_text())
    for d in manifest:
        vid = d["version_id"]
        content = reviewer.get(f"/api/v1/versions/{vid}/content").json()
        assert len(content["text"]) > 3500
        assert hashlib.sha256(content["text"].encode()).hexdigest() == content["source_hash"]
        assert (
            reviewer.get(f"/api/v1/versions/{vid}/download").content
            == (root / d["file"]).read_bytes()
        )
    assert employee.get("/api/v1/documents/ENT-LEGAL").status_code == 404
    assert employee.get("/api/v1/versions/ENT-LEGAL-2.2/content").status_code == 404
    assert employee.get("/api/v1/versions/ENT-LEGAL-2.2/download").status_code == 404
    assert employee.get("/api/v1/versions/ENT-LEGAL-2.2/pages/1").status_code == 404
    for year, value in [("2025-06-01", 18), ("2026-06-01", 20)]:
        q = ask(employee, as_of=year)
        cite = next(c for c in q["citations"] if c["document_id"] == "ENT-LEAVE")
        source = employee.get(f"/api/v1/queries/{q['query_id']}/citations/{cite['id']}").json()
        assert source["text"][cite["start_char"] : cite["end_char"]] == cite["quote"]
        assert cite["reviewed_value"] == value and cite["start_char"] > 1000
        focused = employee.get(
            f"/api/v1/queries/{q['query_id']}/citations/{cite['id']}?view=passage"
        ).json()
        base = focused["text_start_char"]
        assert focused["text"] == source["text"][base : base + len(focused["text"])]
        assert focused["text"][cite["start_char"] - base : cite["end_char"] - base] == cite["quote"]
        assert focused["source_hash"] == source["source_hash"] and len(focused["text"]) < len(
            source["text"]
        )
    conflict = ask(employee, "How many remote work days per week?")
    assert conflict["reason_code"] == "UNRESOLVED_CONFLICT" and conflict["status"] == "abstained"
    assert {c["reviewed_value"] for c in conflict["citations"]} == {2, 3}
    q = ask(employee)
    cite = next(c for c in q["citations"] if c["document_id"] == "ENT-LEAVE")
    with app.state.engine.begin() as conn:
        execute(conn, "DELETE FROM document_acl WHERE document_id='ENT-LEAVE'")
    assert (
        employee.get(f"/api/v1/queries/{q['query_id']}/citations/{cite['id']}").status_code == 404
    )
    assert employee.get(f"/api/v1/versions/{cite['version_id']}/content").status_code == 404
    assert employee.get(f"/api/v1/versions/{cite['version_id']}/download").status_code == 404
    assert (
        ask(employee, "What are our supplier confidentiality obligations?")["status"] == "abstained"
    )


def test_original_pdf_rendering_is_optional_and_permission_checked(system):
    app, _ = system
    add_enterprise(app.state.engine, app.state.settings)
    reviewer = client_for(system, "ravi")
    assert reviewer.get("/api/v1/versions/ENT-LEGAL-2.2/pages/1").status_code == 501
    renderer = os.environ.get("CORTEX_TEST_PDF_RENDERER")
    if renderer:
        configured = create_app(replace(app.state.settings, pdf_renderer=renderer))
        render_client = client_for((configured, system[1]), "ravi")
        response = render_client.get("/api/v1/versions/ENT-LEGAL-2.2/pages/1")
        assert response.status_code == 200 and response.content.startswith(b"\x89PNG\r\n\x1a\n")
        assert int(response.headers["X-PDF-Page-Count"]) >= 2
        assert response.headers["Cache-Control"] == "no-store"
        assert render_client.get("/api/v1/versions/ENT-LEGAL-2.2/pages/201").status_code == 404
        configured.state.engine.dispose()
