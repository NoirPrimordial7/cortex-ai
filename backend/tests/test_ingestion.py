import hashlib
import io
import zipfile
import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from conftest import ask, client_for
from app.security import execute

POLICY = "India full-time employees receive 90 calendar days of notice."


def payload_for(kind):
    output = io.BytesIO()
    if kind == "txt":
        return ("Notice policy\n" + POLICY + "\n").encode()
    if kind == "docx":
        doc = Document()
        doc.add_paragraph("Notice policy")
        doc.add_paragraph(POLICY)
        doc.save(output)
    else:
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 50 700 Td ({POLICY}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.write(output)
    return output.getvalue()


def review_body(content):
    start = content["text"].index(POLICY)
    return dict(
        expected_metadata_revision=0,
        approval_state="approved",
        valid_from="2026-01-01",
        valid_to=None,
        open_ended=True,
        population="india_full_time",
        jurisdiction="IN",
        topic="notice",
        source_kind="hr_policy",
        published_at="2025-12-01",
        start_char=start,
        end_char=start + len(POLICY),
        claim=dict(
            subject="india_full_time",
            predicate="notice_period_days",
            value=90,
            unit="calendar_days",
        ),
        access_scope="uniform",
        reason="Fictional reviewer verifies full source and applicability",
    )


@pytest.mark.parametrize("kind", ["txt", "pdf", "docx"])
def test_actual_upload_private_review_and_citation(system, kind):
    # Isolate this input from the pre-existing seeded notice policy, in this test DB only.
    with system[0].state.engine.begin() as conn:
        execute(
            conn, "DELETE FROM document_acl WHERE tenant_id='NORTHSTAR' AND document_id='HR-NOTICE'"
        )
    reviewer = client_for(system, "ravi")
    employee = client_for(system)
    raw = payload_for(kind)
    response = reviewer.post(
        "/api/v1/documents",
        data={"title": f"Fictional notice {kind}"},
        files={"file": (f"notice.{kind}", raw)},
    )
    assert response.status_code == 202, response.text
    uploaded = response.json()
    assert uploaded["state"] == "pending_review"
    vid, did = uploaded["version_id"], uploaded["document_id"]
    assert employee.get(f"/api/v1/versions/{vid}/content").status_code == 404
    assert ask(reviewer, "What is my notice period?")["status"] == "abstained"
    source = reviewer.get(f"/api/v1/versions/{vid}/content").json()
    assert POLICY in source["text"]
    checked = reviewer.post(f"/api/v1/versions/{vid}/review", json=review_body(source))
    assert checked.status_code == 200, checked.text
    assert checked.json()["state"] == "indexed"
    q = ask(reviewer, "What is my notice period?")
    assert q["status"] == "answered" and q["citations"][0]["reviewed_value"] == 90
    cite = q["citations"][0]
    evidence = reviewer.get(f"/api/v1/queries/{q['query_id']}/citations/{cite['id']}").json()
    assert evidence["text"][cite["start_char"] : cite["end_char"]] == POLICY
    assert evidence["source_hash"] == hashlib.sha256(source["text"].encode()).hexdigest()
    assert reviewer.get(f"/api/v1/versions/{vid}/download").content == raw
    assert ask(employee, "What is my notice period?")["status"] == "abstained"
    acl = reviewer.get(f"/api/v1/admin/documents/{did}/acl").json()
    granted = reviewer.put(
        f"/api/v1/admin/documents/{did}/acl",
        json={
            "expected_policy_revision": acl["policy_revision"],
            "grants": [{"user_id": "ravi"}, {"role_id": "employee"}],
        },
    )
    assert granted.status_code == 200
    assert ask(employee, "What is my notice period?")["status"] == "answered"
    # Publication metadata and original bytes remain separate and immutable.
    detail = reviewer.get(f"/api/v1/documents/{did}").json()["versions"][0]
    assert (
        detail["published_at"] == "2025-12-01"
        and detail["sha256"] == hashlib.sha256(raw).hexdigest()
    )
    assert (
        reviewer.post(f"/api/v1/versions/{vid}/review", json=review_body(source)).status_code == 409
    )


@pytest.mark.parametrize(
    "name,raw",
    [("invalid.txt", b"\xff\x00"), ("fake.pdf", b"NOT A PDF"), ("macro.docx", b"not a zip")],
)
def test_failed_parsing_is_private_and_never_indexed(system, name, raw):
    client = client_for(system, "ravi")
    response = client.post(
        "/api/v1/documents", data={"title": "Rejected input"}, files={"file": (name, raw)}
    )
    assert response.status_code == 202
    result = response.json()
    assert result["state"] == "failed"
    job = client.get(f"/api/v1/ingestion-jobs/{result['job_id']}")
    assert job.status_code == 200 and "EXTRACTION_FAILED_OR_UNSUPPORTED" in job.text
    assert all(
        c["version_id"] != result["version_id"]
        for c in ask(client, "What is my notice period?")["citations"]
    )


def test_metadata_cannot_be_injected_from_upload_fields(system):
    client = client_for(system, "ravi")
    response = client.post(
        "/api/v1/documents",
        data={"title": "Injected approval", "approval_state": "approved", "tenant_id": "ORBIT"},
        files={"file": ("a.txt", b"90 calendar days")},
    )
    assert response.status_code == 422
    assert "Injected approval" not in client.get("/api/v1/documents").text


def test_bounded_archive_and_evidence_span(system):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for index in range(1001):
            archive.writestr(f"part{index}", "")
    client = client_for(system, "ravi")
    response = client.post(
        "/api/v1/documents",
        data={"title": "Too many archive parts"},
        files={"file": ("a.docx", output.getvalue())},
    )
    assert response.json()["state"] == "failed"
    uploaded = client.post(
        "/api/v1/documents",
        data={"title": "Bounded quote"},
        files={"file": ("a.txt", payload_for("txt"))},
    ).json()
    source = client.get(f"/api/v1/versions/{uploaded['version_id']}/content").json()
    body = review_body(source)
    body["end_char"] = body["start_char"] + 8001
    assert (
        client.post(f"/api/v1/versions/{uploaded['version_id']}/review", json=body).status_code
        == 422
    )
